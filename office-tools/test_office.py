from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import office
from docx import Document
from openpyxl import Workbook, load_workbook


class OfficeToolTests(unittest.TestCase):
    def test_docx_replace_crosses_runs_and_includes_tables_and_headers(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.docx"
            output = root / "edited.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("Old ").bold = True
            paragraph.add_run("name")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Old name"
            document.sections[0].header.paragraphs[0].text = "Old name"
            document.save(source)

            result = office.replace_docx(source, output, "Old name", "New name", False)

            self.assertEqual(result["replacements"], 3)
            edited = Document(output)
            self.assertEqual(edited.paragraphs[0].text, "New name")
            self.assertTrue(edited.paragraphs[0].runs[0].bold)
            self.assertEqual(edited.tables[0].cell(0, 0).text, "New name")
            self.assertEqual(edited.sections[0].header.paragraphs[0].text, "New name")

    def test_docx_refuses_implicit_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "document.docx"
            Document().save(path)
            with self.assertRaisesRegex(ValueError, "Refusing to overwrite"):
                office.replace_docx(path, path, "a", "b", False)

    def test_xlsx_set_and_inspect_preserve_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            output = root / "edited.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet["A1"] = "item"
            workbook.save(source)
            workbook.close()

            result = office.set_xlsx(
                source,
                output,
                "Data",
                [["B2", "42"], ["C2", "=B2*2"], ["D2", '"text"']],
                False,
            )

            self.assertTrue(result["verified"])
            edited = load_workbook(output, data_only=False)
            self.assertEqual(edited["Data"]["B2"].value, 42)
            self.assertEqual(edited["Data"]["C2"].value, "=B2*2")
            self.assertEqual(edited["Data"]["D2"].value, "text")
            edited.close()
            inspected = office.inspect_xlsx(output, "Data", "B2:D2", 10, 10, False)
            values = [cell["value"] for cell in inspected["sheets"][0]["rows"][0]]
            self.assertEqual(values, [42, "=B2*2", "text"])


if __name__ == "__main__":
    unittest.main()
