from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries


def print_json(value: Any) -> None:
    print(json.dumps(value, ensure_ascii=False, indent=2, default=json_value))


def json_value(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return value.total_seconds()
    return str(value)


def require_suffix(path: Path, suffixes: set[str]) -> None:
    if path.suffix.lower() not in suffixes:
        expected = ", ".join(sorted(suffixes))
        raise ValueError(f"Unsupported file type {path.suffix!r}; expected {expected}")


def require_input(path: Path, suffixes: set[str]) -> None:
    require_suffix(path, suffixes)
    if not path.is_file():
        raise FileNotFoundError(path)


def ensure_output(input_path: Path | None, output_path: Path, overwrite: bool) -> None:
    if input_path and input_path.resolve() == output_path.resolve() and not overwrite:
        raise ValueError(
            "Refusing to overwrite the original; use another output path or --overwrite"
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)


def iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_document_paragraphs(document: DocumentObject) -> Iterable[Paragraph]:
    seen: set[Any] = set()

    def unique(items: Iterable[Paragraph]) -> Iterable[Paragraph]:
        for paragraph in items:
            marker = paragraph._p
            if marker not in seen:
                seen.add(marker)
                yield paragraph

    yield from unique(document.paragraphs)
    for table in document.tables:
        yield from unique(iter_table_paragraphs(table))
    for section in document.sections:
        for part in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            yield from unique(part.paragraphs)
            for table in part.tables:
                yield from unique(iter_table_paragraphs(table))


def table_values(table: Table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def inspect_docx(path: Path, max_items: int) -> dict[str, Any]:
    require_input(path, {".docx"})
    document = Document(path)
    paragraphs = [
        {
            "index": index,
            "style": paragraph.style.name if paragraph.style else None,
            "text": paragraph.text,
        }
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip()
    ]
    tables = [
        {
            "index": index,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "values": table_values(table),
        }
        for index, table in enumerate(document.tables)
    ]
    return {
        "type": "docx",
        "path": str(path.resolve()),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "paragraphs": paragraphs[:max_items],
        "tables": tables[:max_items],
        "truncated": len(paragraphs) > max_items or len(tables) > max_items,
    }


def worksheet_values(
    worksheet: Any, cell_range: str | None, max_rows: int, max_columns: int
) -> dict[str, Any]:
    if cell_range:
        min_column, min_row, max_column, max_row = range_boundaries(cell_range)
    else:
        min_column = min_row = 1
        max_column = max(worksheet.max_column, 1)
        max_row = max(worksheet.max_row, 1)
    clipped_max_row = min(max_row, min_row + max_rows - 1)
    clipped_max_column = min(max_column, min_column + max_columns - 1)
    rows = []
    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=clipped_max_row,
        min_col=min_column,
        max_col=clipped_max_column,
    ):
        rows.append(
            [
                {
                    "cell": cell.coordinate,
                    "value": cell.value,
                    "data_type": cell.data_type,
                }
                for cell in row
            ]
        )
    return {
        "title": worksheet.title,
        "dimensions": worksheet.calculate_dimension(),
        "requested_range": cell_range,
        "rows": rows,
        "truncated": clipped_max_row < max_row or clipped_max_column < max_column,
    }


def inspect_xlsx(
    path: Path,
    sheet: str | None,
    cell_range: str | None,
    max_rows: int,
    max_columns: int,
    data_only: bool,
) -> dict[str, Any]:
    require_input(path, {".xlsx", ".xlsm"})
    workbook = load_workbook(
        path, data_only=data_only, keep_vba=path.suffix.lower() == ".xlsm"
    )
    try:
        names = [sheet] if sheet else workbook.sheetnames
        missing = [name for name in names if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"Unknown worksheet: {missing[0]}")
        sheets = [
            worksheet_values(workbook[name], cell_range, max_rows, max_columns)
            for name in names
        ]
        return {
            "type": path.suffix.lower().lstrip("."),
            "path": str(path.resolve()),
            "sheet_names": workbook.sheetnames,
            "formulas": not data_only,
            "sheets": sheets,
        }
    finally:
        workbook.close()


def replace_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    runs = paragraph.runs
    cursor = 0
    start_run = start_offset = end_run = end_offset = None
    for index, run in enumerate(runs):
        next_cursor = cursor + len(run.text)
        if start_run is None and start < next_cursor:
            start_run, start_offset = index, start - cursor
        if end <= next_cursor:
            end_run, end_offset = index, end - cursor
            break
        cursor = next_cursor
    if None in (start_run, start_offset, end_run, end_offset):
        return
    first = runs[start_run]
    last = runs[end_run]
    prefix = first.text[:start_offset]
    suffix = last.text[end_offset:]
    first.text = prefix + replacement + (suffix if start_run == end_run else "")
    for index in range(start_run + 1, end_run):
        runs[index].text = ""
    if end_run != start_run:
        last.text = suffix


def replace_paragraph(paragraph: Paragraph, old: str, new: str) -> int:
    text = "".join(run.text for run in paragraph.runs)
    starts: list[int] = []
    cursor = 0
    while True:
        found = text.find(old, cursor)
        if found < 0:
            break
        starts.append(found)
        cursor = found + len(old)
    for start in reversed(starts):
        replace_span(paragraph, start, start + len(old), new)
    return len(starts)


def replace_docx(
    input_path: Path, output_path: Path, old: str, new: str, overwrite: bool
) -> dict[str, Any]:
    require_input(input_path, {".docx"})
    require_suffix(output_path, {".docx"})
    if not old:
        raise ValueError("--old must not be empty")
    ensure_output(input_path, output_path, overwrite)
    document = Document(input_path)
    count = sum(
        replace_paragraph(paragraph, old, new)
        for paragraph in iter_document_paragraphs(document)
    )
    document.save(output_path)
    Document(output_path)
    return {
        "output": str(output_path.resolve()),
        "replacements": count,
        "verified": True,
    }


def parse_cell_value(raw: str) -> Any:
    if raw.startswith("="):
        return raw
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return raw


def set_xlsx(
    input_path: Path,
    output_path: Path,
    sheet: str,
    assignments: list[list[str]],
    overwrite: bool,
) -> dict[str, Any]:
    require_input(input_path, {".xlsx", ".xlsm"})
    require_suffix(output_path, {input_path.suffix.lower()})
    ensure_output(input_path, output_path, overwrite)
    workbook = load_workbook(input_path, keep_vba=input_path.suffix.lower() == ".xlsm")
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Unknown worksheet: {sheet}")
        worksheet = workbook[sheet]
        changed = []
        for cell, raw_value in assignments:
            value = parse_cell_value(raw_value)
            worksheet[cell] = value
            changed.append({"cell": cell, "value": value})
        workbook.save(output_path)
    finally:
        workbook.close()
    check = load_workbook(
        output_path,
        read_only=True,
        data_only=False,
        keep_vba=output_path.suffix.lower() == ".xlsm",
    )
    try:
        verified = [
            {"cell": item["cell"], "value": check[sheet][item["cell"]].value}
            for item in changed
        ]
    finally:
        check.close()
    return {"output": str(output_path.resolve()), "changed": verified, "verified": True}


def new_docx(
    output_path: Path, title: str | None, paragraphs: list[str], overwrite: bool
) -> dict[str, Any]:
    require_suffix(output_path, {".docx"})
    if output_path.exists() and not overwrite:
        raise ValueError("Output already exists; use another path or --overwrite")
    ensure_output(None, output_path, overwrite)
    document = Document()
    if title:
        document.add_heading(title, level=0)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(output_path)
    Document(output_path)
    return {"output": str(output_path.resolve()), "verified": True}


def new_xlsx(
    output_path: Path, sheet: str, rows: list[str], overwrite: bool
) -> dict[str, Any]:
    require_suffix(output_path, {".xlsx"})
    if output_path.exists() and not overwrite:
        raise ValueError("Output already exists; use another path or --overwrite")
    ensure_output(None, output_path, overwrite)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for raw_row in rows:
        row = json.loads(raw_row)
        if not isinstance(row, list):
            raise TypeError("Every --row value must be a JSON array")
        worksheet.append(row)
    workbook.save(output_path)
    workbook.close()
    check = load_workbook(output_path, read_only=True)
    check.close()
    return {"output": str(output_path.resolve()), "verified": True}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Safe local helpers for DOCX and XLSX files"
    )
    commands = parser.add_subparsers(dest="command", required=True)

    inspect_parser = commands.add_parser(
        "inspect", help="print document structure and values as JSON"
    )
    inspect_parser.add_argument("path", type=Path)
    inspect_parser.add_argument("--sheet")
    inspect_parser.add_argument("--range", dest="cell_range")
    inspect_parser.add_argument("--max-rows", type=int, default=100)
    inspect_parser.add_argument("--max-columns", type=int, default=30)
    inspect_parser.add_argument("--max-items", type=int, default=200)
    inspect_parser.add_argument(
        "--data-only", action="store_true", help="show cached formula results"
    )

    replace_parser = commands.add_parser(
        "replace-docx", help="replace text while retaining surrounding styles"
    )
    replace_parser.add_argument("input", type=Path)
    replace_parser.add_argument("output", type=Path)
    replace_parser.add_argument("--old", required=True)
    replace_parser.add_argument("--new", required=True)
    replace_parser.add_argument("--overwrite", action="store_true")

    set_parser = commands.add_parser("set-xlsx", help="set one or more cell values")
    set_parser.add_argument("input", type=Path)
    set_parser.add_argument("output", type=Path)
    set_parser.add_argument("--sheet", required=True)
    set_parser.add_argument(
        "--set",
        dest="assignments",
        action="append",
        nargs=2,
        required=True,
        metavar=("CELL", "VALUE"),
    )
    set_parser.add_argument("--overwrite", action="store_true")

    docx_parser = commands.add_parser("new-docx", help="create a simple DOCX")
    docx_parser.add_argument("output", type=Path)
    docx_parser.add_argument("--title")
    docx_parser.add_argument("--paragraph", action="append", default=[])
    docx_parser.add_argument("--overwrite", action="store_true")

    xlsx_parser = commands.add_parser(
        "new-xlsx", help="create a simple XLSX from JSON rows"
    )
    xlsx_parser.add_argument("output", type=Path)
    xlsx_parser.add_argument("--sheet", default="Sheet1")
    xlsx_parser.add_argument("--row", action="append", default=[])
    xlsx_parser.add_argument("--overwrite", action="store_true")
    return parser


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    args = build_parser().parse_args()
    try:
        if args.command == "inspect":
            suffix = args.path.suffix.lower()
            if suffix == ".docx":
                result = inspect_docx(args.path, args.max_items)
            elif suffix in {".xlsx", ".xlsm"}:
                result = inspect_xlsx(
                    args.path,
                    args.sheet,
                    args.cell_range,
                    args.max_rows,
                    args.max_columns,
                    args.data_only,
                )
            else:
                raise ValueError("inspect supports .docx, .xlsx, and .xlsm")
        elif args.command == "replace-docx":
            result = replace_docx(
                args.input, args.output, args.old, args.new, args.overwrite
            )
        elif args.command == "set-xlsx":
            result = set_xlsx(
                args.input, args.output, args.sheet, args.assignments, args.overwrite
            )
        elif args.command == "new-docx":
            result = new_docx(args.output, args.title, args.paragraph, args.overwrite)
        else:
            result = new_xlsx(args.output, args.sheet, args.row, args.overwrite)
        print_json(result)
        return 0
    except Exception as error:  # noqa: BLE001 - CLI must return machine-readable errors
        print_json({"error": str(error), "type": type(error).__name__})
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
